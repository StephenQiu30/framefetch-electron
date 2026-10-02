from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .files import MAX_DOCUMENT

MAX_TEXT = 500000


def extract_document(path: Path) -> str:
    if path.stat().st_size > MAX_DOCUMENT:
        raise ValueError("document_too_large")
    extension = path.suffix.lower()
    if extension in {".txt", ".md", ".fountain"}:
        data = path.read_bytes()
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("document_encoding_unsupported") from exc
    elif extension == ".pdf":
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ValueError("protected_document")
        if len(reader.pages) > 1000:
            raise ValueError("document_too_many_pages")
        chunks = []
        count = 0
        for page in reader.pages:
            chunk = page.extract_text() or ""
            count += len(chunk)
            if count > MAX_TEXT:
                raise ValueError("document_text_too_large")
            chunks.append(chunk)
        text = "\n\n".join(chunks)
    elif extension == ".docx":
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 10000 or sum(item.file_size for item in members) > 100 * 1024**2:
                raise ValueError("document_expansion_too_large")
            if any(item.file_size > max(item.compress_size, 1) * 200 for item in members):
                raise ValueError("document_expansion_too_large")
        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(
                "\t".join(cell.text for cell in row.cells) for row in table.rows
            )
    else:
        raise ValueError("document_format_unsupported")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()
    if not text:
        raise ValueError("document_has_no_text")
    if len(text) > MAX_TEXT:
        raise ValueError("document_text_too_large")
    return text


def export_docx(markdown: str, destination: Path) -> None:
    document = Document()
    for line in markdown.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        else:
            document.add_paragraph(line)
    document.save(str(destination))
